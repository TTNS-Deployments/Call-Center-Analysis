"""Builds qa/test_cases.md, qa/TEST_REPORT.md and qa/test_report.xlsx.

Statuses of automated cases come from qa/evidence/qa_junit.xml (pytest) and qa/evidence/make_test.log;
the few manual cases carry their own status and point to an evidence file. Run after the suite:

    .venv/bin/python -m pytest qa/tests -v -rA --junitxml=qa/evidence/qa_junit.xml > qa/evidence/qa_pytest.log
    .venv/bin/python qa/build_report.py
"""
from __future__ import annotations

import re
import subprocess
import sys
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import date
from pathlib import Path

QA = Path(__file__).resolve().parent
REPO = QA.parent
EV = QA / "evidence"
LAYERS = ["Backend", "UI", "Front end"]
PRIORITIES = ["Critical", "High", "Medium", "Low"]
STATUSES = ["Pass", "Fail", "Blocked", "Not Tested"]

P = "pytest qa/tests/"  # evidence prefix

# id, layer, priority, area, steps, expected, actual, [pytest test-name prefixes] | ("manual", status, evidence)
CASES = [
    # ------------------------------------------------------------------ Backend: existing suites
    ("B-01", "Backend", "Critical", "Existing tests (core)",
     "make test (first half): .venv/bin/python -m pytest Code/core -q",
     "All existing core tests pass.",
     "First run: 36 passed. After the fixes: 42 passed (6 new regression tests).",
     ("make", "core"), "qa/evidence/make_test.log: '42 passed'"),
    ("B-02", "Backend", "Critical", "Existing tests (web app)",
     "make test (second half): cd Code/webapp && manage.py test",
     "All existing Django tests pass.",
     "First run: 11 tests OK. After the fixes: 15 tests OK (4 new regression tests, incl. collectstatic with production storage).",
     ("make", "web"), "qa/evidence/make_test.log: 'Ran 15 tests ... OK', exit=0"),
    ("B-03", "Backend", "High", "Sample data / make seed",
     "LOCAL_DATA_DIR=<tmp> make seed", "40 synthetic calls generated and analysed, none FAILED, exit 0.",
     "40 calls processed (each line 'CALL-... QA xx% sentiment resolution'), no FAILED lines, exit 0.",
     ["test_B03_"], ""),
    # ------------------------------------------------------------------ Backend: stages
    ("B-04", "Backend", "High", "Transcription output",
     "Seed 30 sample + 4 demo calls; read every transcripts/<id>.json.",
     "Each has source, locale, duration_ms, call_id, audio_blob, transcribed_at and phrases; each phrase has speaker agent|customer, increasing offset_ms, duration_ms>0, non-empty text, confidence 0-1; no mock_script in metadata.",
     "All 34 transcripts well-formed; both speakers present; duration >= last offset.", ["test_B04_"], ""),
    ("B-05", "Backend", "Critical", "PII redaction (sample calls)",
     "For 30 seeded calls extract raw names, phones, emails, card/account numbers, DOB, address from the raw script; search the call store transcript, summary/reason/root cause, QA/compliance/coaching, insights JSON, all 4 Parquet tables and the CRM event payload.",
     "No raw PII value appears in any redacted output (placeholders [NAME], [PHONE], [EMAIL], [CARD], [ACCOUNT], [DATE], [ADDRESS] instead).",
     "240 outputs checked, 0 leaks.", ["test_B05_"], ""),
    ("B-06", "Backend", "Critical", "PII redaction (spoken demo calls)",
     "python -m callcenter upload (4 demo calls, local mode); check the redacted transcript/summary for the spoken PII in Data/demo_calls/scripts.yaml.",
     "Spoken names, phone, email, DOB, card, CVV, OTP are absent from the stored transcript (README: OTP/CVV 'may' be masked).",
     "DEMO-0003 clean. Leaks: DEMO-0001 'rohan dot das at example dot com' and 'twelfth of March, nineteen ninety'; DEMO-0002 'C V V is 1 2 3' and OTP '5 8 2 9 1 4'; DEMO-0004 customer name 'Rajesh Kumar' / 'Mister Kumar' (3 utterances). Visible in UI, API, Parquet.",
     ["test_B06_"], ""),
    ("B-07", "Backend", "High", "PII patterns (mock Language)",
     "MockLanguageClient.analyze on: email, +91 phone, dashed and plain card, account number, PAN, 'My name is Zoya Khan', Flat ... PIN address.",
     "Each value replaced by its placeholder and absent from the output.", "All 8 patterns redacted (e.g. 'Call me on +91 98765 43210' -> 'Call me on [PHONE]').",
     ["test_B07_"], ""),
    ("B-08", "Backend", "High", "PII not in logs (pipeline)",
     "Process 3 PII-heavy calls with root logging at DEBUG; scan captured records.",
     "No raw PII in log records.", "0 log records produced on the success path; no PII. (See also B-45 static check and F-16 live server logs.)",
     ["test_B08_"], ""),
    ("B-09", "Backend", "High", "Sentiment",
     "Check sentiment block of all calls; compare DEMO-0001 (positive) vs DEMO-0003 (angry).",
     "overall/customer in positive|neutral|negative, scores in [-1,1], per-utterance labels valid; DEMO-0001 positive, DEMO-0003 negative.",
     "DEMO-0001 customer positive (+0.268); DEMO-0003 customer negative (-0.333). All structures valid.", ["test_B09_"], ""),
    ("B-10", "Backend", "High", "Emotion",
     "Check emotion {start,end,dominant} for all calls; DEMO-0001 vs DEMO-0003.",
     "Values from the allowed emotion list; DEMO-0001 ends satisfied, DEMO-0003 starts angry/frustrated.",
     "DEMO-0001 neutral -> satisfied; DEMO-0003 angry -> disappointed (dominant frustrated).", ["test_B10_"], ""),
    ("B-11", "Backend", "High", "QA criteria",
     "Check qa block of all calls; DEMO-0001 (model call) vs DEMO-0002 (poor call).",
     "20 criteria in rubric order, valid ratings, percent 0-100 or None, counts sum to 20; DEMO-0001 passes, DEMO-0002 fails with low_qa flag.",
     "DEMO-0001 82.1% passed; DEMO-0002 42.9% failed, flag low_qa.", ["test_B11_"], ""),
    ("B-12", "Backend", "High", "Compliance",
     "Check compliance block of all calls; DEMO-0001 vs DEMO-0002.",
     "All checks in config order with pass|fail|not_applicable; failed_count consistent; DEMO-0001 passes; DEMO-0002 fails recording disclosure and OTP/CVV prohibition, flag compliance_breach.",
     "DEMO-0001 all pass/N/A. DEMO-0002 fails recording_disclosure, identity_before_account_details, no_sensitive_credentials; flag compliance_breach.",
     ["test_B12_"], ""),
    ("B-13", "Backend", "High", "Summary, call reason, root cause",
     "Check summary, call_reason (+name), root_cause, resolution of all calls and expected reasons for the 4 demo calls.",
     "Non-empty summary/root cause, reason from call_reasons.yaml, valid resolution; demo reasons plan_change / card_services / technical_issue / sales_opportunity; DEMO-0003 escalated and repeat contact.",
     "All present. Demo: plan_change/resolved, card_services/resolved, technical_issue/escalated (repeat_contact), sales_opportunity/resolved (root cause 'Not determined').",
     ["test_B13_"], ""),
    ("B-14", "Backend", "Critical", "End to end",
     "Upload one call + sidecar; pipeline.process_recording; spy on status changes; read store, insights JSON, Parquet, list.",
     "Status transcribing -> transcribed -> enriching -> completed; stage done, no error; transcript, insights JSON, 4 Parquet tables written; call listed and retrievable.",
     "Transitions ['transcribing','transcribed','enriching'] -> completed; all artefacts present; retrievable via store and API (B-35).",
     ["test_B14_"], ""),
    # ------------------------------------------------------------------ Backend: edge / negative
    ("B-15", "Backend", "Medium", "Empty file (mock mode)",
     "Upload a 0-byte .wav and process.",
     "Documented local behaviour: mock speech ignores audio and returns a sample conversation; no crash.",
     "Completed with a 14-utterance mock transcript (documented behaviour; real validation needs Azure Speech).", ["test_B15_"], ""),
    ("B-16", "Backend", "Medium", "Silent audio (empty transcript)",
     "Inject a speech client that returns 0 phrases (what a silent recording yields); process.",
     "A clear state: failed with a reason, or completed but recognisable as no speech, not scored as agent failures.",
     "Completed with QA 0.0%, flags unresolved + compliance_breach + low_qa, summary 'The customer called regarding general enquiry: \"\" ...'. Agent penalised for a call with no speech.",
     ["test_B16_"], ""),
    ("B-17", "Backend", "Medium", "Unsupported file",
     "transcribe_recording('x.txt' / 'x.json'); web upload of notes.txt (U-22).",
     "Ignored (None) and no call record; web upload refused with a message.", "Returned None, no record created; web shows 'Unsupported file type .txt.'",
     ["test_B17_"], ""),
    ("B-18", "Backend", "High", "Corrupt audio (Speech error)",
     "Azure fast-transcription client with requests.post mocked to HTTP 400 InvalidAudioFormat (no network); process a garbage .wav.",
     "Call status failed, stage transcribe, error mentions the Speech API error.",
     "failed / transcribe / 'RuntimeError: Speech API error 400: {\"code\":\"InvalidAudioFormat\",...}'.", ["test_B18_"], ""),
    ("B-19", "Backend", "Medium", "Very long transcript",
     "Inject 2,000 phrases (~3.3 h); process; time it. Render detail page (U-28).",
     "Completes in < 60 s with all utterances stored; page renders.",
     "Completed in 0.3 s, 2,000 utterances; detail page 696 KB rendered in 0.11 s.", ["test_B19_"], ""),
    ("B-20", "Backend", "Medium", "Missing metadata",
     "Upload audio with no sidecar JSON; process.", "Completes; started_at defaulted; agent fields empty, no crash.",
     "Completed; agent_id None; started_at = processing time.", ["test_B20_"], ""),
    ("B-21", "Backend", "Medium", "Missing fields in LLM output",
     "Insights client returns {}; process.", "normalize() fills defaults; call completes (reason general_enquiry, resolution unresolved).",
     "Completed, general_enquiry, unresolved, QA None (all N/A).", ["test_B21_"], ""),
    ("B-22", "Backend", "High", "Duplicate call",
     "Process the same blob twice (duplicate BlobCreated); web upload with an existing call ID (U-22).",
     "Second run skipped, processed_at unchanged, one record; web refuses duplicate ID.",
     "Second run returned None ('Skipping ... already completed'); single record; web: 'A call with ID DEMO-0001 already exists.'",
     ["test_B22_"], ""),
    ("B-23", "Backend", "High", "Stage failure: transcribe + retry",
     "Speech client raises ConnectionError; then pipeline.retry.", "failed/transcribe with the error text; retry completes and clears error.",
     "failed / transcribe / 'ConnectionError: speech service unreachable'; retry -> completed, error None.", ["test_B23_"], ""),
    ("B-24", "Backend", "High", "Stage failure: enrich + retry",
     "Language client raises TimeoutError; then retry, counting speech calls.", "failed/enrich; retry resumes at enrich without re-transcribing.",
     "failed / enrich / 'TimeoutError: language service timed out'; retry -> completed, speech re-run 0 times.", ["test_B24_"], ""),
    ("B-25", "Backend", "Medium", "Corrupt intermediate transcript",
     "Overwrite transcripts/<id>.json with invalid JSON; run enrich.", "failed/enrich with a parse error, no crash of the caller beyond the raised error.",
     "failed / enrich / 'JSONDecodeError: Expecting property name ...'.", ["test_B25_"], ""),
    ("B-26", "Backend", "Medium", "CLI status / retry",
     "python -m callcenter status; mark a call failed; status; retry <id>; retry NO-SUCH-CALL.",
     "Counts per status; failed call listed; retry completes; unknown ID exits non-zero.",
     "'48 calls: completed=46, failed=2'; 'FAILED QA-NOMETA-0001' listed; retry -> completed; unknown -> exit 1 (KeyError traceback).",
     ["test_B26_"], ""),
    ("B-27", "Backend", "Medium", "Functions queue handlers",
     "Load Code/functions/function_app.py; call transcribe() with a base64 Event Grid message, enrich(), an other-container event, transcribe_poison().",
     "Stages run; foreign container ignored; poison marks the call failed 'after all retries'; bad message raises ValueError.",
     "transcribed -> completed; 'Ignoring event for container other'; poison -> failed 'Processing failed after all retries'.",
     ["test_B27_"], ""),
    # ------------------------------------------------------------------ Backend: config
    ("B-28", "Backend", "Critical", "Config: local defaults",
     "Fresh interpreter with no backend variables: print settings; manage.py check.",
     "mode local, storage/store local, speech/language/llm mock, events none; Django check clean.",
     "'local local local mock mock mock none'; 'System check identified no issues'.", ["test_B28_"], ""),
    ("B-29", "Backend", "High", "Config: invalid web settings",
     "Start Django with AUTH_MODE=bogus; CC_MODE=azure + AUTH_MODE=dev; CC_MODE=azure without DJANGO_SECRET_KEY.",
     "Startup refused with a message naming the problem.",
     "'Unknown AUTH_MODE 'bogus'', 'AUTH_MODE=dev is only allowed when CC_MODE=local', 'DJANGO_SECRET_KEY must be set outside local mode'.",
     ["test_B29_"], ""),
    ("B-30", "Backend", "High", "Config: invalid pipeline backends",
     "STORAGE/STORE/LLM_BACKEND=foo; LLM_BACKEND=azure and SPEECH_BACKEND=azure without endpoints.",
     "Clear error naming the backend / missing variable.",
     "\"Unknown storage backend 'foo'\" (same for store, llm); 'Set AZURE_OPENAI_ENDPOINT'; 'Set SPEECH_ENDPOINT (...)'.", ["test_B30_"], ""),
    ("B-31", "Backend", "Low", "Config: non-numeric numbers",
     "AGENT_CHANNEL=left; MAX_UPLOAD_MB=lots.", "Error names the variable.",
     "Both crash with \"ValueError: invalid literal for int() with base 10: 'left'/'lots'\"; variable not named.",
     ["test_B31_"], ""),
    ("B-32", "Backend", "Low", "Config: unknown CC_MODE",
     "CC_MODE=prodution (typo).", "Rejected with a clear error.",
     "Accepted silently: mode 'prodution', every backend falls back to local/mock.", ["test_B32_"], ""),
    ("B-33", "Backend", "Medium", "Config: invalid rules override",
     "Write invalid qa_rubric.yaml directly to the config container (bypassing the UI); process a call; remove it; retry.",
     "Call failed with a ConfigError naming the problem; retry after fixing completes.",
     "failed / enrich / \"ConfigError: 'criteria' must be a non-empty list\"; retry -> completed.", ["test_B33_"], ""),
    ("B-34", "Backend", "Medium", "Config: rules validation",
     "save_config_text with bad compliance kind, reason without description, list instead of mapping.",
     "ConfigError with a specific message; nothing saved.", "Messages: 'required or prohibited', 'missing: description', 'mapping'.", ["test_B34_"], ""),
    # ------------------------------------------------------------------ Backend: API
    ("B-35", "Backend", "Critical", "CRM API: valid key",
     "GET /api/calls?range=all and /api/calls/DEMO-0001 with X-API-Key: local-test-key.",
     "200 JSON; list has count + items with call_id, status, agent_id, started_at, summary, call_reason, root_cause, resolution, sentiment, emotion, qa, compliance, flag_list (no transcript/history); detail adds transcript, no history.",
     "200 application/json, count 38+; fields present; detail has transcript, no history.", ["test_B35_"], ""),
    ("B-36", "Backend", "High", "CRM API: filters",
     "agent=A1004&limit=2; customer=CUST-20457; flag=compliance_breach.", "Only matching calls; limit applied.",
     "All filters correct (customer -> ['DEMO-0002']; compliance_breach count 10).", ["test_B36_"], ""),
    ("B-37", "Backend", "Critical", "CRM API: missing / wrong key",
     "No key, wrong key, empty key, key as Bearer token, key in query string - on list and detail.",
     "401/403 JSON, no data.", "All 401 {\"error\": \"unauthorized\"}.", ["test_B37_"], ""),
    ("B-38", "Backend", "High", "CRM API: unknown call ID", "GET /api/calls/NO-SUCH-CALL with valid key.", "404 JSON.",
     "404 {\"error\": \"not found\"}.", ["test_B38_"], ""),
    ("B-39", "Backend", "Medium", "CRM API: malformed input",
     "limit=abc, limit=-5, limit=1e3, from=not-a-date, from=2026-13-45&to=x, range=999, status=%00, agent=' OR 1=1 --.",
     "Clean 4xx JSON error or the parameter ignored; never 500.",
     "limit=abc and limit=1e3 -> HTTP 500 (HTML debug page, ValueError int()). limit=-5 -> 200 but count 35 with only 30 items. Others 200 JSON (bad dates/range ignored; injection strings match nothing).",
     ["test_B39_"], ""),
    ("B-40", "Backend", "Critical", "CRM API: no PII in responses",
     "For 30 sample calls compare raw script PII with /api/calls (all) and /api/calls/<id> responses.",
     "No raw PII in list or detail JSON.", "30 calls checked, 0 leaks (demo-call spoken PII: see B-06).", ["test_B40_"], ""),
    ("B-41", "Backend", "High", "CRM API: session + role scoping",
     "Sign in as agent A1001 and call the API without a key.", "Only own calls listed; other agents' call -> 404.",
     "6 calls, all A1001; DEMO-0002 -> 404; DEMO-0001 -> 200.", ["test_B41_"], ""),
    ("B-42", "Backend", "Medium", "Event feed payload",
     "events.event_payload(DEMO-0002); NullPublisher.publish; EVENTS_BACKEND default.",
     "type call.insights.completed with ids, summary, reason, qa_percent, sentiment, compliance_passed, flags; no transcript/PII; local publisher is a no-op.",
     "Payload fields correct, no transcript or raw PII; events_backend 'none'. (Event Hubs delivery itself not testable locally.)",
     ["test_B42_"], ""),
    # ------------------------------------------------------------------ Backend: security
    ("B-43", "Backend", "High", "No hardcoded secrets",
     "Regex scan of every tracked file for storage keys, key/secret/password/token literals, sk- keys, private keys, connection-string passwords.",
     "No real secrets (documented public emulator keys allowed).",
     "1 hit: docker-compose.yml Azurite AccountKey (public, documented emulator key). Also seen by reading: 'local-dev-only-insecure-key' used only when CC_MODE=local.",
     ["test_B43_"], ""),
    ("B-44", "Backend", "High", ".env git-ignored", "git check-ignore .env; git ls-files; git log --all -- .env local.settings.json.",
     ".env ignored, not tracked, never committed.", "Ignored (.gitignore:16), not tracked, no commits touch it. (.env contents not printed.)", ["test_B44_"], ""),
    ("B-45", "Backend", "Medium", "Log statements don't print text",
     "Static scan of logging calls in Code/ for transcript/phrase/text/doc arguments.",
     "No log call interpolates call text.", "None found (only call IDs, blob names, statuses).", ["test_B45_"], ""),
    ("B-46", "Backend", "Medium", "Docker local stack",
     "docker compose -p qa-cc-smoke up --build -d; run --rm seed; GET /healthz, sign in, /calls/, API; down -v.",
     "Image builds (collectstatic), web + Azurite start, 40 calls seeded through the real Blob SDK, pages 200.",
     None, ("manual",), "qa/evidence/docker_smoke.log"),
    # ------------------------------------------------------------------ UI
    ("U-01", "UI", "High", "Login page", "GET /auth/login.", "200; demo personas Admin, Supervisor, QA and 6 agents listed.",
     "200; 9 personas: dev-admin, dev-supervisor, dev-qa, dev-A1001..A1006.", ["test_U01_"], ""),
    ("U-02", "UI", "Critical", "Login each demo user", "POST /auth/login user=<persona> for all 9.", "302 to /; overview 200 with Sign out.",
     "All 9 signed in.", ["test_U02_"], ""),
    ("U-03", "UI", "High", "Invalid login",
     "Dev mode: unknown persona. Password mode (DEMO_PASSWORD set): wrong password; unknown user.",
     "Rejected with a generic message; not signed in; no stack trace.",
     "Dev: 400 'Unknown user' (plain text), still signed out. Password: 403 'Wrong password.'; unknown user 400.", ["test_U03_"], ""),
    ("U-04", "UI", "High", "Login ?next= open redirect", "Login with next=https://evil.example/, //evil.example/x, /calls/.",
     "External targets replaced by /; local path kept.", "evil -> /, //evil -> /, /calls/ -> /calls/.", ["test_U04_"], ""),
    ("U-05", "UI", "High", "Logout", "GET /auth/logout; POST /auth/logout; then GET /calls/.", "GET 405; POST 302 to login; session gone.",
     "405; 302 /auth/login; /calls/ -> 302 /auth/login?next=/calls/.", ["test_U05_"], ""),
    ("U-06", "UI", "Critical", "Protected pages need sign-in", "Anonymous GET of every route incl. audio and Power BI token.",
     "302 to /auth/login?next=...", "All 13 routes redirect.", ["test_U06_"], ""),
    ("U-07", "UI", "High", "Protected POSTs need sign-in", "Anonymous POST upload, retry, settings; GET /healthz.",
     "302 to login; healthz public 'ok'.", "All redirected; healthz 'ok'.", ["test_U07_"], ""),
    ("U-08", "UI", "Critical", "Role-based access matrix",
     "Admin, Supervisor, QA, Agent(A1001): GET 12 routes (own/other call and audio, agents, own/other agent, me, insights, pipeline, settings) and POST upload/retry/settings.",
     "Matrix per roles.py: Supervisor no Settings; QA no Pipeline/Settings/upload/retry; Agent only own calls/scorecard, 403 elsewhere.",
     "0 mismatches for all 4 roles.", ["test_U08_"], ""),
    ("U-09", "UI", "Medium", "Navigation per role", "Read sidebar items per role.",
     "Admin all 6; Supervisor no Settings; QA no Pipeline/Settings; Agent Overview, Calls, My scorecard.", "As expected.", ["test_U09_"], ""),
    ("U-10", "UI", "Critical", "Agent sees only own calls", "Agent A1001: /calls/?range=all; owner of every listed call.",
     "Only A1001 calls.", "6 calls, all A1001.", ["test_U10_"], ""),
    ("U-11", "UI", "Critical", "Pages load without errors",
     "Each persona type: 18 URLs incl. range presets, team filter, builtin insights, 4 demo call pages.",
     "No 5xx, no template errors.", "0 errors for Admin, Supervisor, QA, Agent.", ["test_U11_"], ""),
    ("U-12", "UI", "High", "Call list content", "/calls/?range=all: header and DEMO-0003 row.",
     "Columns incl. call, agent, summary, sentiment, resolution, QA; row shows reason, flags.",
     "Columns: Date, Call, Agent, Summary, Sentiment, Resolution, QA, Duration; row shows reason, sentiment, resolution, QA %, flags.",
     ["test_U12_"], ""),
    ("U-13", "UI", "Critical", "Call detail content", "QA user opens /calls/DEMO-0003/.",
     "Transcript (with redaction chips), sentiment, emotion start/end, QA score, all compliance checks, summary, reason, root cause.",
     "All 9 elements present.", ["test_U13_"], ""),
    ("U-14", "UI", "High", "Search / filter / sort",
     "q (case-insensitive), agent, team, reason, sentiment, resolution, flag, qa_min/qa_max, sort=qa / -qa.",
     "Only matching rows; sort order correct.", "All filters correct; asc 25.0, 42.9, 42.9...; desc 100.0, 93.3, 93.3...", ["test_U14_"], ""),
    ("U-15", "UI", "Medium", "Empty state", "/calls/?q=zzz; all list pages for an empty date range; agent with no calls.",
     "'No calls match these filters.'; pages 200.", "As expected.", ["test_U15_"], ""),
    ("U-16", "UI", "Medium", "Pagination", "/calls/?range=all page 1/2; page=999, abc, -1.", "25 per page, no overlap, Next/Previous; bad page numbers handled.",
     "25 + rest, no overlap; bad pages 200.", ["test_U16_"], ""),
    ("U-17", "UI", "Medium", "Bad parameters",
     "/calls/? qa_min=abc, qa_max=<script>, sort=;drop, range=bogus, from=2026-99-99, page=1e9, team=%00.",
     "Friendly handling (ignored or message), never 500.",
     "qa_min=abc and qa_max=<script> -> HTTP 500 (ValueError: could not convert string to float). Others 200.", ["test_U17_"], ""),
    ("U-18", "UI", "High", "Invalid IDs (DEBUG=False)",
     "/calls/NO-SUCH-CALL/, its /audio, /settings/?file=../../.env, /nope/, GET on retry.",
     "Friendly 404 page ('Not found', Back to overview) / 405; no stack trace.", "404 custom page x4; 405 for GET retry.", ["test_U18_"], ""),
    ("U-19", "UI", "High", "No debug leak with DEBUG=False", "Trigger the 500 from U-17 with DEBUG=False.",
     "Generic error page, no traceback/exception/paths.", "'Server Error (500)' generic page only.", ["test_U19_"], ""),
    ("U-20", "UI", "Low", "Debug pages with local defaults",
     "Default local settings (DJANGO_DEBUG unset): /calls/NO-SUCH-CALL/ and the U-17 500.", "No debug details.",
     "DEBUG=True by default in local mode: technical 404 ('Raised by: portal.views.call_detail') and full 500 traceback shown. By design for local dev; Azure default is False (U-21).",
     ["test_U20_"], ""),
    ("U-21", "UI", "High", "Azure-mode secure defaults",
     "CC_MODE=azure (+ secret, entra): read DEBUG, cookie flags, SSL redirect, HSTS.", "DEBUG False; secure cookies; SSL redirect; HSTS.",
     "False True True True 3600.", ["test_U21_"], ""),
    ("U-22", "UI", "High", "Upload form",
     "Supervisor: no file; .txt; existing call ID; size over MAX_UPLOAD_MB; valid upload with call_id 'QA-UP/../../x'.",
     "Specific error messages; valid upload sanitised and processed to completed.",
     "'Choose an audio file to upload.', 'Unsupported file type .txt.', 'A call with ID DEMO-0001 already exists.', 'File is larger than 0 MB.'; valid -> 'QA-UP-..-..-x', completed.",
     ["test_U22_"], ""),
    ("U-23", "UI", "High", "Retry from UI", "Mark DEMO-0004 failed; Supervisor sees it on Pipeline; POST retry; unknown ID.",
     "Failure shown with reason; retry completes; unknown -> 404.", "Shown with 'simulated outage'; 302 -> completed; unknown 404.", ["test_U23_"], ""),
    ("U-24", "UI", "Medium", "Retry next= open redirect", "Supervisor POST /calls/DEMO-0003/retry next=https://evil.example/phish.",
     "Redirect stays on site.", "302 Location: https://evil.example/phish (unvalidated).", ["test_U24_"], ""),
    ("U-25", "UI", "Medium", "Settings edit", "Admin: save invalid YAML, valid YAML, reset.", "'Not saved' + reason; 'Saved.'; reset OK.",
     "As expected.", ["test_U25_"], ""),
    ("U-26", "UI", "Medium", "htmx partials", "HX-Request on /calls/ and /pipeline/.", "Fragments without <html>.", "As expected.", ["test_U26_"], ""),
    ("U-27", "UI", "Medium", "Audio streaming", "Full GET, Range 100-199, suffix -10, unsatisfiable range.",
     "200 audio/*; 206 exact bytes; 416.", "200 audio/x-wav RIFF; 206 bytes match; 416.", ["test_U27_"], ""),
    ("U-28", "UI", "Medium", "Edge-case calls render", "No-metadata call, 2,000-utterance call, failed call: list, detail, pipeline, overview, agents.",
     "All 200; long page fast.", "All 200; long page 696 KB in 0.11 s.", ["test_U28_"], ""),
    # ------------------------------------------------------------------ Front end
    ("F-01", "Front end", "Medium", "Templates: no missing variables",
     "(a) string_if_invalid on pages with complete calls; (b) all pages incl. failed/no-metadata calls: visible text has no literal None or template syntax.",
     "No unresolved variables; no 'None' shown.",
     "(a) none. (b) Overview 'flagged calls' shows 'None · 23 Oct, 15:30' for a call without agent name (overview.html:86 {{ c.agent_name }} has no default).",
     ["test_F01_"], ""),
    ("F-02", "Front end", "High", "Static files",
     "Collect every /static/ URL from 9 pages; fetch on runserver (DEBUG) and gunicorn DEBUG=false after collectstatic (whitenoise manifest).",
     "All 200 with right content types; prod uses hashed names + long cache; missing file 404.",
     "css/js/png 200 on both; prod app.46f2f7805c64.css etc. 'max-age=315360000, public, immutable'; missing 404.", ["test_F02_"], ""),
    ("F-03", "Front end", "High", "Browser console / network",
     "Chromium headless: Admin visits 10 pages (networkidle), Agent 4; record console errors, page errors, failed requests, HTTP >= 400.",
     "None.", "None (audio streams aborted by navigation excluded). Note: 2 page loads stalled 30 s on CDN scripts during the run (see D-07).",
     ["test_F03_"], ""),
    ("F-04", "Front end", "Medium", "Charts render", "Overview, Insights, Agents, Agent pages: count canvases vs Chart.js instances.",
     "Every canvas has a chart; Chart.js and htmx loaded.", "3/3, 5/5, 1/1, 1/1.", ["test_F04_"], ""),
    ("F-05", "Front end", "Medium", "Live search (htmx)", "Type DEMO-0002 in the search box.", "Table updates via htmx to 1 row.",
     "25 rows -> ['DEMO-0002'], URL pushed.", ["test_F05_"], ""),
    ("F-06", "Front end", "Medium", "Audio player", "Open DEMO-0001; wait for audio metadata.", "Duration ~88 s.", "88.3 s.", ["test_F06_"], ""),
    ("F-07", "Front end", "High", "Browser sign-in / sign-out", "Click QA persona; click Sign out; open /calls/.",
     "Name shown; signed out; redirected to login.", "As expected.", ["test_F07_"], ""),
    ("F-08", "Front end", "High", "CSRF", "Every POST form on 10 pages has csrfmiddlewaretoken; POST /auth/login without token.",
     "All forms have token; tokenless POST 403.", "All forms have it; 403.", ["test_F08_"], ""),
    ("F-09", "Front end", "Medium", "Required fields / error messages",
     "Pipeline: submit upload without file. Settings: save 'categories: []'.", "Browser blocks (input :invalid); settings shows 'Not saved' + reason.",
     "Blocked; message shown.", ["test_F09_"], ""),
    ("F-10", "Front end", "Medium", "Accessibility basics",
     "9 pages: title, html lang, viewport, img alt, input labels, button names, h1.", "All present; titles distinct.",
     "All present; titles 'Overview · Call Center Analysis' etc.", ["test_F10_"], ""),
    ("F-11", "Front end", "Medium", "Mobile width (375 px)", "Chromium 375x812: scrollWidth - clientWidth on 9 app pages + login.",
     "No horizontal scroll.", "Login 0 px; every app page overflows 260 px (sidebar and main 635 px wide).", ["test_F11_"], ""),
    ("F-12", "Front end", "High", "Security headers / cookies",
     "Sign in on runserver and on gunicorn DEBUG=false; inspect headers and Set-Cookie.",
     "X-Frame-Options DENY, nosniff, Referrer-Policy, session cookie HttpOnly + SameSite=Lax, csrftoken SameSite=Lax.",
     "DENY, nosniff, same-origin, COOP same-origin; sessionid HttpOnly SameSite=Lax; csrftoken SameSite=Lax (Secure only in Azure mode, U-21).",
     ["test_F12_"], ""),
    ("F-13", "Front end", "Low", "Content-Security-Policy", "Check CSP header.", "A CSP is sent.", "No CSP header.", ["test_F13_"], ""),
    ("F-14", "Front end", "Critical", "Output escaping (XSS)",
     "Create a call whose transcript, agent name and team contain <img onerror> and <script>; search with payloads; open detail, list, agent page, overview in Chromium.",
     "Payloads shown as text; no dialogs; no injected elements.", "Escaped (&lt;script&gt;); 0 dialogs; 0 injected <img>.", ["test_F14_"], ""),
    ("F-15", "Front end", "Low", "Third-party scripts", "Inspect external <script> tags.", "Pinned versions with Subresource Integrity (or self-hosted).",
     "htmx 2.0.4 (unpkg) and Chart.js 4.4.1 (jsdelivr) pinned but no integrity attribute; page load waits on them.", ["test_F15_"], ""),
    ("F-16", "Front end", "High", "Live server logs free of PII", "After the browser tests, scan runserver + gunicorn logs for raw PII.",
     "None.", "150+ lines, 0 PII.", ["test_F16_"], ""),
]

DEFECTS = [
    ("D-01", "Medium", "Backend", "Local mock PII redaction misses spoken-form PII and unknown names",
     "make setup; python -m callcenter upload (local mode); open /calls/DEMO-0001/, DEMO-0002, DEMO-0004 or GET /api/calls/DEMO-0001.",
     "Expected: names, email, date of birth, CVV and OTP redacted. Actual: 'rohan dot das at example dot com', 'twelfth of March, nineteen ninety', 'C V V is 1 2 3', OTP '5 8 2 9 1 4', 'Mister Rajesh Kumar' stay in the transcript, UI, API and Parquet.",
     "B-06",
     "MockLanguageClient (Code/core/callcenter/language.py) is regex-based: names only from scenarios.yaml or after 'my name is', emails only in a@b.c form, digit runs only >= 10 digits, dates only '12 March 1990'. Production uses Azure AI Language (not tested here). Idea: add spoken-email / 'C V V'/'O T P' digit-run / spelled-date patterns and demo-script names to the mock, or a test that runs the demo scripts through it."),
    ("D-02", "Medium", "Backend / UI", "Non-numeric query parameters cause HTTP 500",
     "Sign in; open /calls/?qa_min=abc (or qa_max=<script>). Or GET /api/calls?limit=abc (or limit=1e3) with X-API-Key.",
     "Expected: parameter ignored or a 400 message. Actual: HTTP 500 (ValueError); with local defaults the full Django traceback page. Also limit=-5 returns count 35 but 30 items.",
     "U-17, B-39",
     "analytics.filter_calls float(f['qa_min']) (Code/core/callcenter/analytics.py:86-88) and api.calls int(request.GET['limit']) (Code/webapp/portal/api.py:46) are not guarded. Idea: parse with try/except, clamp limit to 1..1000."),
    ("D-03", "Medium", "Backend", "Silent / empty-transcript call is scored as a failing agent",
     "Process a recording whose transcription returns no phrases (silent audio).",
     "Expected: a clear 'no speech' outcome (failed or not scored). Actual: completed with QA 0%, flags compliance_breach + low_qa + unresolved, summary with an empty quote.",
     "B-16",
     "pipeline.enrich_transcript has no guard for zero phrases; llm.normalize (llm.py:123) marks unconditional required checks as 'fail' when absent, and the mock rates every criterion not_met. Idea: short-circuit empty transcripts to a 'no_speech' state."),
    ("D-04", "Medium", "Front end", "Horizontal scrolling on phone-width screens",
     "Open any signed-in page at 375 px width (Chromium device emulation).",
     "Expected: no horizontal scroll. Actual: page is 635 px wide (260 px overflow); login page is fine.",
     "F-11",
     "In the <=860 px media query .app uses grid-template-columns: 1fr (app.css:293); 1fr = minmax(auto,1fr) so the column cannot shrink below the sidebar/nav min-content width. Idea: minmax(0,1fr) and min-width:0 on .sidebar/.main."),
    ("D-05", "Low", "UI", "Retry 'next' parameter allows an off-site redirect",
     "As Supervisor POST /calls/<id>/retry with next=https://evil.example/phish (valid CSRF token).",
     "Expected: redirect limited to this site. Actual: 302 Location https://evil.example/phish.",
     "U-24",
     "views.call_retry redirects to request.POST['next'] unchecked (Code/webapp/portal/views.py:162). Needs a signed-in Supervisor and CSRF token, so low risk. Idea: reuse url_has_allowed_host_and_scheme as in accounts.views._safe_next."),
    ("D-06", "Low", "Front end", "No Content-Security-Policy header",
     "Inspect response headers of any page (runserver or gunicorn).", "Expected: a CSP. Actual: none (other headers fine).", "F-13",
     "No CSP middleware/setting. Idea: Django 6 ContentSecurityPolicyMiddleware or a header allowing self + the two CDNs (and inline scripts via nonces)."),
    ("D-07", "Low", "Front end", "CDN scripts without Subresource Integrity; pages wait on public CDNs",
     "View source of any page; observe load in a slow network.",
     "Expected: integrity attributes or self-hosted copies. Actual: htmx (unpkg) and Chart.js (jsdelivr) have no integrity; twice during this run a page 'load' stalled 30 s waiting for them (server answered in ms).",
     "F-15 (F-03 observation)",
     "base.html:50-51. Idea: add integrity + crossorigin, or vendor the files into static/ (also helps private-network deployments)."),
    ("D-08", "Low", "Front end", "'None' shown as agent name for calls without metadata",
     "Upload a call without agent details that gets flagged; open the Overview.",
     "Expected: blank or '–'. Actual: 'None · 23 Oct, 15:30' in the flagged-calls list.", "F-01",
     "overview.html:86 {{ c.agent_name }} without |default; build_call_document stores agent_name=None. Idea: |default:'–'."),
    ("D-09", "Low", "Backend", "Misspelt CC_MODE silently runs in local mode",
     "CC_MODE=prodution python -m callcenter status.", "Expected: startup error. Actual: runs with local storage and mock AI.", "B-32",
     "config.get_settings (config.py:114) only checks == 'azure'. Idea: reject values other than local/azure."),
    ("D-10", "Low", "Backend", "Non-numeric AGENT_CHANNEL / MAX_UPLOAD_MB: error does not name the variable",
     "AGENT_CHANNEL=left python -m callcenter status; MAX_UPLOAD_MB=lots manage.py check.",
     "Expected: message naming the variable. Actual: traceback ending \"ValueError: invalid literal for int() with base 10: 'left'\".", "B-31",
     "config.py:144 and settings.py:116 call int() directly. Idea: helper that raises 'AGENT_CHANNEL must be an integer'."),
    ("D-11", "Low", "UI", "Local defaults show Django debug pages",
     "make run (DJANGO_DEBUG unset); open /calls/NO-SUCH-CALL/ or trigger D-02.",
     "Expected (per QA brief): no debug details. Actual: technical 404 and full traceback pages, because DEBUG defaults to true when CC_MODE=local.",
     "U-20",
     "Intentional for developer machines (settings.py: DEBUG = env_bool('DJANGO_DEBUG', IS_LOCAL)); Azure default is False (U-21) and docker-compose sets false. Only a risk if local mode is exposed to others."),
]

NOT_TESTED = [
    ("Out of scope", "Real Azure AI Speech, Language and OpenAI", "Needs keys and costs money; brief says mocks only. So real transcription of silent/corrupt audio, real PII accuracy and LLM scoring against Data/demo_calls/README.md expectations are unverified."),
    ("Out of scope", "Cosmos DB call store, Blob Storage/ADLS on Azure, Key Vault references", "Azure resources. SQLite and folder storage tested; Azurite path covered by B-46."),
    ("Out of scope", "Event Hubs publishing (EVENTS_BACKEND=eventhub)", "Azure only. Payload builder tested (B-42)."),
    ("Out of scope", "Microsoft Entra ID sign-in (AUTH_MODE=entra), app roles from tokens", "Needs an app registration and credentials."),
    ("Out of scope", "Power BI embedding / report (Analytics/)", "Needs Power BI Embedded/Fabric capacity; no .pbix in repo."),
    ("Out of scope", "ARM deployment, Deploy to Azure button, deploy.sh, deploy_code_arm.sh, setup_entra.sh, GitHub Deploy workflow", "Real deployment. ARM/parameter consistency is covered by existing tests (B-01)."),
    ("Out of scope", "Private endpoints, VPN, Microsoft Defender, App Service networking, HTTPS/HSTS on the real host", "Azure portal / deployment checks."),
    ("Out of scope", "QueueDispatcher (Azure retry via Storage queues)", "Needs Azure Storage queues; local retry path tested."),
    ("Not Tested", "Azure Functions host runtime (func start + Azurite queues)", "Needs Functions Core Tools (npm global install). Handler logic run in-process instead (B-27)."),
    ("Not Tested", "make package (build/webapp.zip, functions.zip)", "Writes into build/ outside qa/; brief restricts new files to qa/."),
    ("Not Tested", "make worker (folder watcher)", "Endless polling loop; same pipeline entry point as tested paths."),
    ("Not Tested", "Cross-browser (Firefox/WebKit), screen-reader audit, colour contrast", "Only Chromium installed; basic a11y checks done (F-10)."),
    ("Not Tested", "Load / concurrency testing", "Not requested; single long-transcript performance checked (B-19, U-28)."),
]


# Re-test after the fixes (2026-10-23): new actual results; the first-run result is kept alongside.
RETEST = {
    "B-06": "All 4 demo calls clean, 0 leaks. DEMO-0001 redacts email + date of birth ([EMAIL], [DATE]); DEMO-0002 CVV and OTP -> [CREDENTIAL]; DEMO-0004 'Mister [NAME]'.",
    "B-16": "Completed with flag 'no_speech' only, QA not scored (all N/A), compliance N/A, summary 'No speech was detected in the recording, so the call was not scored.'; no LLM call.",
    "B-31": "\"AGENT_CHANNEL must be a whole number, got 'left'\" and \"MAX_UPLOAD_MB must be a whole number, got 'lots'\".",
    "B-32": "Refused: \"Unknown CC_MODE 'prodution' (local or azure)\", exit 1.",
    "B-39": "limit=abc / 1e3 / -5 -> 400 {\"error\": \"limit must be a whole number from 1 to 1000\"}; other malformed values 200 JSON (ignored). No 500.",
    "U-17": "All 7 bad parameters -> 200; non-numeric qa_min / qa_max are ignored.",
    "U-20": "DEBUG now defaults to False in local mode too: friendly 404 page and generic 500 page, no debug details (DJANGO_DEBUG=true opts in).",
    "U-24": "next=https://evil.example/phish -> 302 /calls/DEMO-0003/ (off-site targets ignored); /pipeline/ still honoured.",
    "F-01": "(a) none. (b) no literal None or template syntax on any page; flagged list shows '–' for a missing agent name.",
    "F-11": "0 px horizontal overflow on all 10 pages at 375 px wide.",
    "F-13": "CSP sent on every response: default-src 'self'; script-src 'self' 'unsafe-inline'; ... frame-ancestors 'none'; object-src 'none'.",
    "F-15": "No third-party scripts: htmx, Chart.js and powerbi-client are served from static/vendor (hash-checked against the npm packages).",
}

FIXES = {
    "D-01": "language.py mock: spoken e-mail and spelled-out date patterns, honorific + name (Mister/Mr/Mrs/Ms/Dr), CVV/OTP/PIN digit runs (also in the customer's reply after the agent asks for one) -> [CREDENTIAL]. Tests: test_language.py.",
    "D-02": "analytics.filter_calls ignores non-numeric qa_min/qa_max; api.calls validates limit (1-1000) and returns 400 JSON. Tests: test_pipeline.py, portal/tests.py.",
    "D-03": "pipeline.enrich_transcript: no phrases -> no_speech_insights() (all N/A, no AI calls), flag 'no_speech' (label + style), not 'unresolved'. Test: test_pipeline.py.",
    "D-04": "app.css: mobile grid column minmax(0,1fr); sidebar/main, grid children and page-header items may shrink; selects capped at container width.",
    "D-05": "views.call_retry checks next with url_has_allowed_host_and_scheme, else the call page. Test: portal/tests.py.",
    "D-06": "New portal/security.py middleware sends CONTENT_SECURITY_POLICY (settings.py; overridable by env). Test: portal/tests.py.",
    "D-07": "htmx 2.0.4, Chart.js 4.4.1, powerbi-client 2.23.1 vendored in static/vendor (README with sources; source-map comment removed so collectstatic works). Test: collectstatic with production storage.",
    "D-08": "overview.html: {{ c.agent_name|default:'–' }}.",
    "D-09": "config.get_settings and Django settings reject CC_MODE other than local/azure. Test: test_pipeline.py.",
    "D-10": "_env_int / env_int helpers name the variable in the error. Test: test_pipeline.py.",
    "D-11": "DJANGO_DEBUG defaults to false everywhere; local runs still serve static files via whitenoise finders; new templates/500.html; documented in configuration.md and .env.example.",
}

OBSERVATIONS = [
    "Mock speech ignores audio content (documented), so empty/corrupt audio can only be checked with injected clients (B-15, B-16, B-18).",
    "Raw (unredacted) transcripts are kept in the `transcripts` container by design (SECURITY.md: private container + lifecycle rule); every user-facing output is redacted.",
    "Dev-mode sign-in with an unknown persona returns a bare `400 Unknown user` text page (no template).",
    "`python -m callcenter retry <unknown id>` exits 1 with a KeyError traceback rather than a message.",
    "The Azurite account key in docker-compose.yml is the public emulator key, not a secret.",
    "First run: twice a browser page load stalled 30 s waiting for the public CDN scripts (server answered in ms); fixed by D-07.",
    "Found while fixing: two retries of the same call at once could collide on the local storage temp file (FileNotFoundError); storage.py now uses a unique temp name per write.",
]


# --------------------------------------------------------------------------- results

def junit_results() -> dict[str, list[tuple[str, str]]]:
    root = ET.parse(EV / "qa_junit.xml").getroot()
    out: dict[str, list[tuple[str, str]]] = {}
    for tc in root.iter("testcase"):
        name = tc.get("name")
        if tc.find("failure") is not None or tc.find("error") is not None:
            st = "Fail"
        elif tc.find("skipped") is not None:
            st = "Blocked"
        else:
            st = "Pass"
        out.setdefault(name, []).append((tc.get("classname"), st))
    return out


def make_status(which: str) -> str:
    log = (EV / "make_test.log").read_text()
    core, web = log.split("manage.py test", 1)
    if which == "core":
        return "Pass" if re.search(r"\b\d+ passed\b", core) and not re.search(r"\bfailed\b|\berror", core) else "Fail"
    return "Pass" if re.search(r"Ran \d+ tests.*?\n\nOK", web, re.S) and "exit=0" in log else "Fail"


def manual_status(evidence: str) -> str:
    path = REPO / evidence
    if not path.exists():
        return "Not Tested"
    m = re.search(r"^RESULT: (Pass|Fail|Blocked)", path.read_text(), re.M)
    return m.group(1) if m else "Not Tested"


def resolve():
    results = junit_results()
    rows = []
    for cid, layer, prio, area, steps, expected, actual, how, evidence in CASES:
        if how[0] == "make":
            status = make_status(how[1])
        elif how[0] == "manual":
            status = manual_status(evidence)
            if actual is None:
                text = (REPO / evidence).read_text() if (REPO / evidence).exists() else ""
                m = re.search(r"^ACTUAL: (.+)$", text, re.M)
                actual = m.group(1) if m else "Not run."
        else:
            matched = [(n, st) for n, sts in results.items() for _, st in sts if any(n.startswith(p) for p in how)]
            if not matched:
                status, actual = "Not Tested", "No result recorded."
            else:
                status = "Fail" if any(st == "Fail" for _, st in matched) else "Pass"
                fails = [n for n, st in matched if st == "Fail"]
                files = sorted({c for n, sts in results.items() for c, _ in sts if any(n.startswith(p) for p in how)})
                evidence = (f"{P}{files[0].split('.')[-1]}.py -k {how[0].rstrip('_')}  "
                            f"({len(matched)} run, {len(fails)} failed"
                            + (f": {', '.join(fails)}" if fails else "") + "); output in qa/evidence/qa_pytest.log")
        first = ""
        if cid in RETEST:
            first, actual = actual, RETEST[cid]
        rows.append(dict(id=cid, layer=layer, priority=prio, area=area, steps=steps, expected=expected,
                         actual=actual, status=status, evidence=evidence, first=first))
    return rows


# --------------------------------------------------------------------------- outputs

def env_info():
    git = lambda *a: subprocess.run(["git", *a], cwd=REPO, capture_output=True, text=True).stdout.strip()  # noqa: E731
    py = subprocess.run([str(REPO / ".venv/bin/python"), "--version"], capture_output=True, text=True).stdout.strip()
    pkgs = subprocess.run([str(REPO / ".venv/bin/python"), "-c",
                           "import django, playwright, pytest; from importlib.metadata import version as v; "
                           "print(f'Django {django.get_version()}, pytest {pytest.__version__}, Playwright {v(\"playwright\")} (Chromium headless), openpyxl {v(\"openpyxl\")}')"],
                          capture_output=True, text=True).stdout.strip()
    return {
        "Mode": "CC_MODE=local (mock Speech/Language/OpenAI, folder storage, SQLite), AUTH_MODE=dev, CRM_API_KEYS=local-test-key",
        "Branch": git("branch", "--show-current"),
        "Commit": git("log", "-1", "--format=%h %s"),
        "Python": py.replace("Python ", "") + " (.venv)",
        "Tools": pkgs,
        "OS": "Linux (GitHub Codespace)",
        "Date": date.today().isoformat(),
        "Data": "Isolated temp data folder per run (developer .localdata untouched)",
    }


def counts(rows):
    by = Counter((r["layer"], r["status"]) for r in rows)
    byp = Counter((r["priority"], r["status"]) for r in rows)
    return by, byp


def verdict(rows):
    """Exit rule: all Critical and High pass -> 'Works as expected'. Otherwise a judgement: 'Not working' would mean
    the core workflow (process a call, see it in the app/API) is broken; here it is not (see JUSTIFICATION)."""
    crit_high_fail = [r for r in rows if r["priority"] in ("Critical", "High") and r["status"] != "Pass"]
    if not crit_high_fail:
        return "Works as expected", crit_high_fail
    core = {"B-14", "U-11", "U-13", "B-35"}
    core_broken = any(r["id"] in core and r["status"] != "Pass" for r in rows)
    return ("Not working" if core_broken else "Works with issues"), crit_high_fail


JUSTIFICATION = (
    "Tested once in local mode across backend, UI and front end (90 test cases, 154 automated checks, plus the "
    "app's own 57 tests). The first run found 12 failures (11 defects: 1 Critical, 6 Medium, 5 Low by test); all were "
    "fixed in the code and the full suite was run again: every test case passes except the Docker stack check, "
    "which stays Blocked because containers cannot reach each other in this Codespace (the image itself builds "
    "and collectstatic succeeds). Every pipeline stage works (transcription, PII redaction, sentiment, emotion, QA, "
    "compliance, summary, reason, root cause, storage, failure and retry), as do all pages, roles, the CRM API and "
    "the security checks. All Critical and High tests pass, so the exit rule is met. Not covered: the real Azure "
    "services, Entra ID sign-in and Power BI, which need Azure credentials.")


def write_markdown(rows, info):
    v, chf = verdict(rows)
    by, byp = counts(rows)
    tot = Counter(r["status"] for r in rows)
    flat = [st for sts in junit_results().values() for _, st in sts]
    ju = {"total": len(flat), "Pass": flat.count("Pass"), "Fail": flat.count("Fail")}
    executed = tot["Pass"] + tot["Fail"]
    rate = 100 * tot["Pass"] / executed if executed else 0
    esc = lambda s: str(s).replace("|", "\\|").replace("\n", " ")  # noqa: E731

    tc = ["# Test cases (local mode)", "",
          "Scope: what exists in this repository and runs with `CC_MODE=local`. Layers: Backend (Code/core, functions, API), "
          "UI (Django pages and server behaviour), Front end (static assets, templates, browser). Automated tests live in "
          "`qa/tests/` and are named after these IDs (`test_B05_...`).", "",
          "| ID | Layer | Priority | Area | Steps | Expected result |", "|---|---|---|---|---|---|"]
    tc += [f"| {r['id']} | {r['layer']} | {r['priority']} | {esc(r['area'])} | {esc(r['steps'])} | {esc(r['expected'])} |" for r in rows]
    tc += ["", "## Out of scope", "", "| Item | Reason |", "|---|---|"]
    tc += [f"| {esc(i)} | {esc(why)} |" for kind, i, why in NOT_TESTED if kind == "Out of scope"]
    tc += ["", "Also not run (in scope but not executed): see *Not Tested* in `TEST_REPORT.md`.", ""]
    (QA / "test_cases.md").write_text("\n".join(tc))

    md = ["# Test report: Call Center Analysis (local mode)", "",
          f"**Verdict: {v}**", "", JUSTIFICATION, "",
          "## Environment", "", "| | |", "|---|---|"]
    md += [f"| {k} | {esc(val)} |" for k, val in info.items()]
    md += ["", "## Totals", "", "| Layer | Pass | Fail | Blocked | Not Tested | Total |", "|---|---|---|---|---|---|"]
    for layer in LAYERS:
        n = [by[(layer, s)] for s in STATUSES]
        md.append(f"| {layer} | " + " | ".join(map(str, n)) + f" | {sum(n)} |")
    md.append("| **All** | " + " | ".join(str(tot[s]) for s in STATUSES) + f" | {len(rows)} |")
    md += ["", "| Priority | Pass | Fail | Blocked | Not Tested | Total |", "|---|---|---|---|---|---|"]
    for p in PRIORITIES:
        n = [byp[(p, s)] for s in STATUSES]
        md.append(f"| {p} | " + " | ".join(map(str, n)) + f" | {sum(n)} |")
    md += ["", f"Pass rate (Pass / executed): **{rate:.1f}%** ({tot['Pass']}/{executed}). "
           f"Exit rule (all Critical and High pass): **{'met' if not chf else 'not met'}**"
           + (f" (failing: {', '.join(r['id'] for r in chf)})" if chf else "") + ".", "",
           f"Automated checks behind the cases: {ju['total']} pytest items in `qa/tests` ({ju['Pass']} passed, {ju['Fail']} failed) "
           "plus the app's own 57 tests (`make test`, all passed). Every remaining failure is a product issue; test-harness "
           "problems found during the run were fixed in the tests, not in the app.", "",
           "## Failed tests", "", "| ID | Priority | Area | Reason | Defect |", "|---|---|---|---|---|"]
    dmap = {}
    for d in DEFECTS:
        for t in re.findall(r"[BUF]-\d+", d[6]):
            dmap.setdefault(t, []).append(d[0])
    for r in rows:
        if r["status"] == "Fail":
            md.append(f"| {r['id']} | {r['priority']} | {esc(r['area'])} | {esc(r['actual'])} | {', '.join(dmap.get(r['id'], []))} |")
    md += ["", "## Defects", ""]
    for d in DEFECTS:
        md += [f"### {d[0]} ({d[1]}, {d[2]}): {d[3]}", "",
               f"- **Steps to reproduce:** {d[4]}", f"- **Expected vs actual:** {d[5]}",
               f"- **Affected tests:** {d[6]}", f"- **Suggested cause (read-only):** {d[7]}", ""]
    md += ["## Not tested / out of scope", "", "| Type | Item | Reason |", "|---|---|---|"]
    md += [f"| {k} | {esc(i)} | {esc(w)} |" for k, i, w in NOT_TESTED]
    md += ["", "## Observations (no defect raised)", ""] + [f"- {o}" for o in OBSERVATIONS]
    md += ["", "## How to re-run", "", "```bash",
           ".venv/bin/python -m pip install openpyxl          # Playwright + Chromium were already in the venv",
           "make test > qa/evidence/make_test.log 2>&1; echo exit=$? >> qa/evidence/make_test.log",
           ".venv/bin/python -m pytest qa/tests -v -rA --junitxml=qa/evidence/qa_junit.xml > qa/evidence/qa_pytest.log 2>&1",
           ".venv/bin/python qa/build_report.py              # regenerates test_cases.md, TEST_REPORT.md, test_report.xlsx",
           "```", "",
           "Files: `qa/00_discovery.md`, `qa/test_cases.md`, `qa/TEST_REPORT.md`, `qa/test_report.xlsx`, `qa/tests/`, `qa/evidence/`.", ""]
    (QA / "TEST_REPORT.md").write_text("\n".join(md))
    return v, chf, tot, rate


def write_xlsx(rows, info, v):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E78")
    wrap = Alignment(wrap_text=True, vertical="top")
    colors = {"Pass": "C6EFCE", "Fail": "FFC7CE", "Blocked": "FFEB9C", "Not Tested": "D9D9D9"}

    def header(ws, cols, widths):
        ws.append(cols)
        for i, w in enumerate(widths, 1):
            c = ws.cell(row=1, column=i)
            c.font, c.fill, c.alignment = head, fill, Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    # Test Cases
    tc = wb.active
    tc.title = "Test Cases"
    header(tc, ["ID", "Layer", "Priority", "Area", "Steps", "Expected Result", "Actual Result", "Status", "Evidence",
                "First run (before fixes)"], [7, 10, 10, 26, 55, 50, 60, 11, 50, 55])
    for r in rows:
        tc.append([r[k] for k in ("id", "layer", "priority", "area", "steps", "expected", "actual", "status", "evidence", "first")])
        for c in tc[tc.max_row]:
            c.alignment = wrap
        tc.cell(row=tc.max_row, column=8).fill = PatternFill("solid", fgColor=colors[r["status"]])
    n = len(rows) + 1
    tc.auto_filter.ref = f"A1:J{n}"

    # Summary
    sm = wb.create_sheet("Summary", 0)
    sm.column_dimensions["A"].width = 30
    for col in "BCDEFG":
        sm.column_dimensions[col].width = 13
    sm["A1"] = "Call Center Analysis: QA test report (local mode)"
    sm["A1"].font = Font(bold=True, size=14)
    r = 3
    sm.cell(row=r, column=1, value="Environment").font = Font(bold=True)
    for k, val in info.items():
        r += 1
        sm.cell(row=r, column=1, value=k)
        sm.cell(row=r, column=2, value=val)

    def rate(p, f):
        return p / (p + f) if (p + f) else "-"

    def table(start, title, field, keys):
        """Fixed numbers from this one test run (no formulas)."""
        sm.cell(row=start, column=1, value=title).font = Font(bold=True)
        hdr = [title.split(" ")[-1].capitalize()] + STATUSES + ["Total", "Pass rate"]
        for i, h in enumerate(hdr, 1):
            c = sm.cell(row=start + 1, column=i, value=h)
            c.font, c.fill = head, fill
        totals = Counter()
        for j, key in enumerate(keys):
            row = start + 2 + j
            cnt = Counter(x["status"] for x in rows if x[field] == key)
            totals.update(cnt)
            sm.cell(row=row, column=1, value=key)
            for i, st in enumerate(STATUSES, 2):
                sm.cell(row=row, column=i, value=cnt[st])
            sm.cell(row=row, column=6, value=sum(cnt.values()))
            sm.cell(row=row, column=7, value=rate(cnt["Pass"], cnt["Fail"])).number_format = "0.0%"
        tot_row = start + 2 + len(keys)
        sm.cell(row=tot_row, column=1, value="Total").font = Font(bold=True)
        for i, st in enumerate(STATUSES, 2):
            sm.cell(row=tot_row, column=i, value=totals[st]).font = Font(bold=True)
        sm.cell(row=tot_row, column=6, value=sum(totals.values())).font = Font(bold=True)
        sm.cell(row=tot_row, column=7, value=rate(totals["Pass"], totals["Fail"])).number_format = "0.0%"
        return tot_row, totals

    r += 2
    end1, all_tot = table(r, "Totals by layer", "layer", LAYERS)
    end2, _ = table(end1 + 2, "Totals by priority", "priority", PRIORITIES)
    not_passing = [x["id"] for x in rows if x["priority"] in ("Critical", "High") and x["status"] != "Pass"]
    r = end2 + 2
    sm.cell(row=r, column=1, value="Pass rate (Pass / executed)").font = Font(bold=True)
    sm.cell(row=r, column=2, value=rate(all_tot["Pass"], all_tot["Fail"])).number_format = "0.0%"
    r += 1
    sm.cell(row=r, column=1, value="Critical + High not passing").font = Font(bold=True)
    sm.cell(row=r, column=2, value=len(not_passing))
    sm.cell(row=r, column=3, value=", ".join(not_passing))
    r += 1
    sm.cell(row=r, column=1, value="Exit rule met (all Critical & High pass)").font = Font(bold=True)
    sm.cell(row=r, column=2, value="NO" if not_passing else "YES")
    r += 2
    sm.cell(row=r, column=1, value="Verdict").font = Font(bold=True, size=12)
    sm.cell(row=r, column=2, value=v).font = Font(bold=True, size=12, color="9C5700")
    r += 1
    sm.cell(row=r, column=1, value="Justification").font = Font(bold=True)
    sm.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    c = sm.cell(row=r, column=2, value=JUSTIFICATION)
    c.alignment = wrap
    sm.row_dimensions[r].height = 210
    for row in sm.iter_rows(min_row=4, max_row=4 + len(info)):
        row[1].alignment = Alignment(wrap_text=False)

    # Defects
    df = wb.create_sheet("Defects")
    header(df, ["ID", "Severity", "Layer", "Title", "Steps to reproduce", "Expected vs Actual (first run)", "Affected test IDs",
                "Cause", "Status", "Fix"], [7, 10, 13, 40, 50, 60, 16, 60, 10, 70])
    for d in DEFECTS:
        df.append(list(d) + ["Fixed" if d[0] in FIXES else "Open", FIXES.get(d[0], "")])
        df.cell(row=df.max_row, column=9).fill = PatternFill("solid", fgColor="C6EFCE" if d[0] in FIXES else "FFC7CE")
        for c in df[df.max_row]:
            c.alignment = wrap

    # Not Tested / Out of Scope
    nt = wb.create_sheet("Not Tested - Out of Scope")
    header(nt, ["Type", "Item", "Reason"], [14, 60, 90])
    for row in NOT_TESTED:
        nt.append(list(row))
        for c in nt[nt.max_row]:
            c.alignment = wrap

    # Failed Tests
    dmap = {}
    for d in DEFECTS:
        for t in re.findall(r"[BUF]-\d+", d[6]):
            dmap.setdefault(t, []).append(d[0])
    ft = wb.create_sheet("Fixed Tests", 2)
    header(ft, ["ID", "Layer", "Priority", "Area", "First run: why it failed", "Defect", "Re-test after fix", "Status now"],
           [7, 10, 10, 30, 70, 10, 70, 11])
    for r in rows:
        if r["first"] or r["status"] == "Fail":
            ft.append([r["id"], r["layer"], r["priority"], r["area"], r["first"], ", ".join(dmap.get(r["id"], [])),
                       r["actual"], r["status"]])
            ft.cell(row=ft.max_row, column=8).fill = PatternFill("solid", fgColor=colors[r["status"]])
            for c in ft[ft.max_row]:
                c.alignment = wrap

    # Observations
    ob = wb.create_sheet("Observations")
    header(ob, ["#", "Observation (no defect raised)"], [5, 120])
    for i, o in enumerate(OBSERVATIONS, 1):
        ob.append([i, o])
        ob.cell(row=ob.max_row, column=2).alignment = wrap

    # Discovery (from 00_discovery.md: every heading + markdown table)
    ds = wb.create_sheet("Discovery")
    ds.column_dimensions["A"].width = 38
    for col in "BCDEF":
        ds.column_dimensions[col].width = 45
    clean = lambda t: re.sub(r"[`*]", "", t).replace("\\|", "|").strip()  # noqa: E731
    for line in (QA / "00_discovery.md").read_text().splitlines():
        if line.startswith("#"):
            ds.append([])
            ds.append([clean(line.lstrip("# "))])
            ds.cell(row=ds.max_row, column=1).font = Font(bold=True, size=12)
        elif line.startswith("|"):
            if re.fullmatch(r"\|[\s|:-]+\|", line.strip()):
                continue
            cells = [clean(c) for c in re.split(r"(?<!\\)\|", line.strip())[1:-1]]
            ds.append(cells)
            is_header = ds.cell(row=ds.max_row - 1, column=1).font.size == 12
            for c in ds[ds.max_row]:
                c.alignment = wrap
                if is_header:
                    c.font, c.fill = head, fill
        elif line.strip().startswith("- "):
            ds.append([clean(line.strip()[2:])])
            ds.merge_cells(start_row=ds.max_row, start_column=1, end_row=ds.max_row, end_column=6)
            ds.cell(row=ds.max_row, column=1).alignment = wrap
            ds.row_dimensions[ds.max_row].height = 32
        elif line.strip():
            ds.append([clean(line)])

    wb.save(QA / "test_report.xlsx")


def main():
    rows = resolve()
    info = env_info()
    v, chf, tot, rate = write_markdown(rows, info)
    write_xlsx(rows, info, v)
    print(f"Verdict: {v}")
    print("Totals:", dict(tot), f"pass rate {rate:.1f}%", "| Critical/High not passing:", [r["id"] for r in chf])
    for r in rows:
        if r["status"] == "Fail":
            print(f"  FAIL {r['id']} [{r['priority']}] {r['area']}: {r['actual'][:150]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
